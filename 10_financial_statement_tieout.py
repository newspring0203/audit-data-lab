"""Mission 10: 가상 2025년 GL → 시산표 → 재무제표 초안 푸팅/대사.
"""
from __future__ import annotations

import argparse
import csv
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parent
DATA = ROOT / 'data'
CASE = DATA / 'mission10'
REPORTS = ROOT / 'reports'
YEAR = '2025'

# 계정별 정상 잔액 방향. 1590은 차감 계정(감가상각누계액).
ACCOUNTS = {
    '1010': ('보통예금', '자산', '차변'),
    '1100': ('매출채권', '자산', '차변'),
    '1200': ('재고자산', '자산', '차변'),
    '1500': ('유형자산', '자산', '차변'),
    '1590': ('감가상각누계액', '자산 차감', '대변'),
    '2100': ('장기차입금', '부채', '대변'),
    '2200': ('미지급비용', '부채', '대변'),
    '3100': ('자본금', '자본', '대변'),
    '4010': ('매출', '수익', '대변'),
    '5010': ('매출원가', '비용', '차변'),
    '6100': ('급여', '비용', '차변'),
    '6200': ('임차료', '비용', '차변'),
    '6300': ('감가상각비', '비용', '차변'),
    '7100': ('이자비용', '비용', '차변'),
}

# 원장 계정의 순액을 재무제표에 표시할 때의 부호: 양수=증가/수익, 음수=차감/비용.
FS_ACCOUNT_MAP = {
    'cash': ('1010', 1),
    'accounts_receivable': ('1100', 1),
    'inventory': ('1200', 1),
    'equipment': ('1500', 1),
    'accumulated_depreciation': ('1590', -1),
    'loan_payable': ('2100', 1),
    'accrued_expenses': ('2200', 1),
    'share_capital': ('3100', 1),
    'revenue': ('4010', 1),
    'cost_of_sales': ('5010', -1),
    'salaries': ('6100', -1),
    'rent': ('6200', -1),
    'depreciation': ('6300', -1),
    'interest_expense': ('7100', -1),
}


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open('r', encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


def write_csv(path: Path, data: list[dict], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction='ignore')
        writer.writeheader()
        writer.writerows(data)


def load_combined_journals():
    existing_h = read_csv(DATA / 'derived/journal_headers.csv')
    existing_l = read_csv(DATA / 'derived/journal_lines.csv')
    extra_h = read_csv(CASE / 'extra_journal_headers.csv')
    extra_l = read_csv(CASE / 'extra_journal_lines.csv')
    headers = existing_h + extra_h
    lines = existing_l + extra_l

    header_by_id = {}
    for h in headers:
        if h['entry_id'] in header_by_id:
            raise ValueError(f'전표 헤더 ID 중복: {h["entry_id"]}')
        header_by_id[h['entry_id']] = h
    seen_lines = set()
    for line in lines:
        k = (line['entry_id'], line['line_no'])
        if k in seen_lines:
            raise ValueError(f'전표 상세 행 ID 중복: {k}')
        seen_lines.add(k)
        if line['entry_id'] not in header_by_id:
            raise ValueError(f'헤더 없는 전표 상세: {k}')
        if line['account_code'] not in ACCOUNTS:
            raise ValueError(f'미등록 계정코드: {line["account_code"]}')
        for amount_field in ('debit_krw', 'credit_krw'):
            amt = int(line[amount_field])
            if amt < 0:
                raise ValueError(f'음수 금액: {k} {amount_field}')
        if int(line['debit_krw']) > 0 and int(line['credit_krw']) > 0:
            raise ValueError(f'한 행에 차변/대변 동시 입력: {k}')

    for h in headers:
        if h['posting_date'][:4] != str(h['fiscal_year']):
            raise ValueError(f'전기일과 회계연도 불일치: {h["entry_id"]}')
    return headers, lines, header_by_id


def trial_balance(headers, lines, header_by_id):
    totals = defaultdict(lambda: {'debit': 0, 'credit': 0})
    by_entry = defaultdict(lambda: {'debit': 0, 'credit': 0})
    for line in lines:
        h = header_by_id[line['entry_id']]
        if h['fiscal_year'] != YEAR or h['posting_status'] != 'POSTED':
            continue  # 2026년 1월 입금은 2025년 재무제표에 넣지 않음
        d = int(line['debit_krw']); c = int(line['credit_krw'])
        totals[line['account_code']]['debit'] += d
        totals[line['account_code']]['credit'] += c
        by_entry[line['entry_id']]['debit'] += d
        by_entry[line['entry_id']]['credit'] += c
    unbalanced = [eid for eid, a in by_entry.items() if a['debit'] != a['credit']]
    if unbalanced:
        raise ValueError(f'차대변 불일치 전표: {unbalanced}')

    rows = []
    for code in sorted(ACCOUNTS):
        name, category, side = ACCOUNTS[code]
        deb = totals[code]['debit']; cre = totals[code]['credit']
        rows.append({
            'account_code': code, 'account_name': name, 'account_type': category,
            'normal_side': side, 'debit_turnover_krw': deb, 'credit_turnover_krw': cre,
            'closing_debit_krw': max(deb-cre, 0),
            'closing_credit_krw': max(cre-deb, 0),
        })
    if sum(r['closing_debit_krw'] for r in rows) != sum(r['closing_credit_krw'] for r in rows):
        raise AssertionError('시산표가 불일치합니다')
    return rows, len(by_entry)


def gl_account_amount(tb_rows, code: str) -> int:
    row = next(r for r in tb_rows if r['account_code'] == code)
    side = ACCOUNTS[code][2]
    if side == '차변':
        return row['closing_debit_krw'] - row['closing_credit_krw']
    return row['closing_credit_krw'] - row['closing_debit_krw']


def expected_fs_lines(tb_rows) -> dict[str, int]:
    values = {
        line_id: sign * gl_account_amount(tb_rows, acc)
        for line_id, (acc, sign) in FS_ACCOUNT_MAP.items()
    }
    values.update({
        'current_assets': values['cash']+values['accounts_receivable']+values['inventory'],
        'net_equipment': values['equipment']+values['accumulated_depreciation'],
    })
    values['total_assets'] = values['current_assets']+values['net_equipment']
    values['total_liabilities'] = values['loan_payable']+values['accrued_expenses']
    values['gross_profit'] = values['revenue']+values['cost_of_sales']
    values['operating_expenses'] = values['salaries']+values['rent']+values['depreciation']
    values['operating_profit'] = values['gross_profit']+values['operating_expenses']
    values['net_income'] = values['operating_profit']+values['interest_expense']
    values['current_year_profit'] = values['net_income']
    values['total_equity'] = values['share_capital']+values['current_year_profit']
    values['total_liabilities_equity'] = values['total_liabilities']+values['total_equity']
    return values


def compare_draft(tb_rows):
    fs = read_csv(CASE / 'draft_financial_statements.csv')
    statement = {}
    for r in fs:
        if r['line_id'] in statement:
            raise ValueError(f'재무제표 초안 항목 ID 중복: {r["line_id"]}')
        statement[r['line_id']] = int(r['reported_krw'])
    gl_values = expected_fs_lines(tb_rows)
    checks = []

    def add(code, category, label, reported, expected, audit_focus, source):
        delta = reported - expected
        checks.append({
            'check_id': code, '검증구분': category, '검증내용': label,
            '보고금액_원': reported, '기대금액_원': expected,
            '차이_원': delta, '결과': '일치' if delta == 0 else '차이 발견',
            '감사상확인사항': audit_focus, '비교기준': source,
        })

    # 1. 푸팅(보고서 안의 더하기 검증). 기준: 제출 초안에 표시된 하위 항목.
    formulas = {
        'current_assets': ['cash','accounts_receivable','inventory'],
        'net_equipment': ['equipment','accumulated_depreciation'],
        'total_assets': ['current_assets','net_equipment'],
        'total_liabilities': ['loan_payable','accrued_expenses'],
        'total_equity': ['share_capital','current_year_profit'],
        'total_liabilities_equity': ['total_liabilities','total_equity'],
        'gross_profit': ['revenue','cost_of_sales'],
        'operating_expenses': ['salaries','rent','depreciation'],
        'operating_profit': ['gross_profit','operating_expenses'],
        'net_income': ['operating_profit','interest_expense'],
    }
    names = {r['line_id']: r['line_name'] for r in fs}
    for out, ins in formulas.items():
        if out not in statement or any(k not in statement for k in ins):
            raise ValueError(f'재무제표 초안 필수 항목 누락: {out} <- {ins}')
        add(f'FOOT_{out.upper()}', '푸팅',
            f'{names[out]} = ' + ' + '.join(names[k] for k in ins),
            statement[out], sum(statement[k] for k in ins),
            '산술 합계 오기·수식 및 보고서 버전 확인', '재무제표 초안 내부 합계')

    # 2. 계정별 GL↔재무제표 연결 및 GL에서 도출한 수익·자본 관련 교차 검증.
    for line_id in FS_ACCOUNT_MAP:
        add(f'GL_{line_id.upper()}', '원장대사', names[line_id],
            statement[line_id], gl_values[line_id],
            '전표 원본·계정 분류·결산조정·작성본 변경 내역 확인', '2025년 전기된 GL / 결산 추가 전표')
    for line_id in ('net_income','current_year_profit'):
        add(f'GL_{line_id.upper()}', '원장대사', names[line_id],
            statement[line_id], gl_values[line_id],
            '손익 집계·보고서 버전·미반영 결산조정 확인', '2025년 GL 산출 순이익')

    # 3. 재무제표 사이 연결과 재무상태표 회계등식.
    add('CROSS_PROFIT', '표간연결', '손익계산서 당기순이익 ↔ 재무상태표 당기이익',
        statement['net_income'], statement['current_year_profit'],
        '두 재무제표의 수정·갱신 시점 및 숫자 참조 확인', '재무제표 초안 간 일치 여부')
    add('BS_EQUATION', '회계등식', '자산총계 = 부채·자본총계',
        statement['total_assets'], statement['total_liabilities_equity'],
        '자산·부채·자본 분류 및 총계 계산 점검', '재무상태표 회계등식')
    return checks, fs, gl_values


def build_workpaper():
    headers, lines, by_id = load_combined_journals()
    tb_rows, entry_count = trial_balance(headers, lines, by_id)
    checks, statement_rows, gl_values = compare_draft(tb_rows)
    exceptions = []
    for r in checks:
        if r['결과'] != '차이 발견':
            continue
        exceptions.append({
            '검토번호': r['check_id'], '검증구분': r['검증구분'],
            '발견사항': r['검증내용'], '초안금액_원': r['보고금액_원'],
            '비교기준금액_원': r['기대금액_원'], '차이_원': r['차이_원'],
            '추가확인': r['감사상확인사항'],
            '감사인검토의견': '', '최종결론': '판단 보류'
        })
    total_debit = sum(r['closing_debit_krw'] for r in tb_rows)
    total_credit = sum(r['closing_credit_krw'] for r in tb_rows)
    summary = [
        {'항목': '분석 회계연도', '값': YEAR},
        {'항목': '2025년 전표 수', '값': entry_count},
        {'항목': '시산표 차변 합계', '값': total_debit},
        {'항목': '시산표 대변 합계', '값': total_credit},
        {'항목': '검증 규칙 수', '값': len(checks)},
        {'항목': '차이 검출 규칙 수', '값': len(exceptions)},
        {'항목': '2025년 GL 매출액', '값': gl_values['revenue']},
        {'항목': '2025년 GL 당기이익(간이)', '값': gl_values['net_income']},
        {'항목': '자료 범위', '값': '기존 가상 매출·입금 GL + 별도 가상 결산·기초 투자 전표'},
        {'항목': '주의', '값': '고의 오류가 들어간 가상 재무제표 초안. 감사의견/오류 확정 아님'},
    ]
    return summary, tb_rows, statement_rows, checks, exceptions


def export_csv(summary, tb, statements, checks, exceptions):
    REPORTS.mkdir(parents=True, exist_ok=True)
    write_csv(REPORTS / 'mission10_trial_balance_2025.csv', tb, list(tb[0]))
    write_csv(REPORTS / 'mission10_draft_statements.csv', statements, list(statements[0]))
    write_csv(REPORTS / 'mission10_checks.csv', checks, list(checks[0]))
    write_csv(REPORTS / 'mission10_workpaper.csv', exceptions, list(exceptions[0]))
    write_csv(REPORTS / 'mission10_summary.csv', summary, ['항목','값'])


def export_excel(summary, tb, statements, checks, exceptions):
    # 사용자는 Mission 07에서 pandas/openpyxl을 설치한 상태이므로 같은 방식을 활용.
    # 이 함수는 CSV 검증 로직과 분리하여 엑셀 엔진 없이도 --csv-only 실행 가능.
    import pandas as pd
    from openpyxl.styles import Font, PatternFill
    out = REPORTS / 'mission10_financial_statement_review.xlsx'
    with pd.ExcelWriter(out, engine='openpyxl') as writer:
        for name, data in [
            ('00_요약', summary), ('01_시산표2025', tb),
            ('02_재무제표초안', statements), ('03_검증상세', checks),
            ('04_감사검토조서', exceptions),
        ]:
            pd.DataFrame(data).to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='17365D')
            for column in ws.columns:
                letter = column[0].column_letter
                max_length = max(len(str(cell.value or '')) for cell in column)
                ws.column_dimensions[letter].width = min(max(max_length + 3, 12), 48)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, (float, int)):
                        cell.number_format = '#,##0;[Red](#,##0);-'
    return out


def main():
    parser = argparse.ArgumentParser(description='Mission 10: GL → 시산표 → 푸팅/대사')
    parser.add_argument('--csv-only', action='store_true', help='CSV만 생성(Excel 엔진 불필요)')
    args = parser.parse_args()
    summary, tb, statements, checks, exceptions = build_workpaper()
    export_csv(summary, tb, statements, checks, exceptions)
    print('=== Mission 10 | 2025년 간이 재무제표 푸팅·원장 대사 ===')
    for row in exceptions:
        print(f"{row['검토번호']}: 초안 {row['초안금액_원']:,} / 기준 {row['비교기준금액_원']:,} / 차이 {row['차이_원']:+,}원")
    print(f'검증 {len(checks)}항목 / 차이 {len(exceptions)}항목 / CSV 저장: {REPORTS}')
    if not args.csv_only:
        print(f'Excel 저장: {export_excel(summary, tb, statements, checks, exceptions)}')
    print('※ 전표 추가분과 재무제표 초안은 교육용 가상자료, 탐지 결과는 확정 감사 결론이 아닙니다.')


if __name__ == '__main__':
    main()
